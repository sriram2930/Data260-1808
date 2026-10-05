"""One-off script assembling the HW5 Word report from the extracted
screenshots plus METRICS.md / AI_USE.md / REFLECTION.md content. Not part of
the graded deliverable, kept here only so the report can be regenerated."""

import docx
from docx.shared import Inches, Pt

SS = "screenshots/"

doc = docx.Document()

doc.add_heading("DATA 260 Homework 5 Report", level=0)
doc.add_paragraph("Sreeram Achutuni")
doc.add_paragraph("SJSU ID: 019151808")
doc.add_paragraph("GitHub repository: https://github.com/sriram2930/Data260-1808")

h1 = lambda t: doc.add_heading(t, level=1)
h2 = lambda t: doc.add_heading(t, level=2)
p = lambda t: doc.add_paragraph(t)

def img(path, width=6.3):
    doc.add_picture(SS + path, width=Inches(width))

def code_block(text):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)

def note(text):
    para = doc.add_paragraph()
    run = para.add_run("Note: " + text)
    run.italic = True

h1("Section 0: Config values")
config_rows = [
    ("SID4", "1808"),
    ("PORT_BASE", "8008"),
    ("PREFIX", "s1808"),
    ("SEED", "1808"),
    ("VERIFY_SEED", "261808"),
    ("DOMAIN_ID", "0, Campus course catalogue and enrolment"),
    ("Hardware", "Intel Core i7-11390H, 16GB RAM, integrated Iris Xe graphics (no dedicated GPU)"),
    ("Local model", "qwen2.5:1.5b-instruct via Ollama, used by the Part 5 agent loop"),
    ("Repo", "https://github.com/sriram2930/Data260-1808"),
    ("Collaborators", "Sbnikitha and supriyaselvanganesan added as collaborators with repo access"),
    ("Commit tagged as hw5", "5256e97f9f998ec526fb2448cd37baf8ddf4b361"),
]
t = doc.add_table(rows=0, cols=2)
t.style = "Table Grid"
for k, v in config_rows:
    row = t.add_row()
    row.cells[0].text = k
    row.cells[1].text = v

# ============================================================= PART 1
h1("Part 1: Database, API, and Redux Client")
p(
    "This builds on HW4's FastAPI and MySQL backend instead of starting over. Department is "
    "the new related entity (name, building, a unique dept_code, timestamps), and Course now has "
    "a department_id foreign key, a seats_available field with a default of 30, timestamps, and a "
    "unique course_code. The migration ran against the already-seeded 5,000-row table in place, "
    "nothing was dropped. The React client's course data moved from component state to Redux "
    "Toolkit, with async thunks over axios calling the same FastAPI endpoints."
)

h2("Department: create")
p("POST /api/departments, creating a Physics department.")
img("p01_img1.jpeg")

h2("Department: list (paginated)")
p("GET /api/departments?skip=0&limit=5, returning the items/total/skip/limit envelope.")
img("p01_img2.png")

h2("Department: get by id")
img("p02_img1.png")

h2("Department: update")
p("PUT /api/departments/1, changing the building field. updated_at changes, created_at does not.")
img("p02_img2.jpeg")

h2("Department: delete blocked (still in use)")
p("DELETE on a department that still has courses pointing at it returns 409, not a raw DB error.")
img("p03_img1.png")

h2("Department: delete allowed (not in use)")
p("Deleting the Physics department created earlier, which has no courses attached, succeeds.")
img("p03_img2.png")

h2("Relationship query: courses in a department")
p("GET /api/departments/1/courses, scoped to Computer Science.")
img("p04_img1.png")

h2("Course validation: bad course_code format")
img("p04_img2.png")

h2("Course validation: duplicate course_code")
img("p05_img1.png")

h2("Database contents")
p("Querying MySQL directly to confirm both tables hold the data the API shows, including the "
  "updated building value from the PUT above.")
img("p05_img2.png", width=6.3)

h2("Redux: Home screen")
p("Home.jsx reads the course list from Redux state with useSelector and dispatches "
  "fetchCourses() itself, instead of receiving the list as a prop from App.jsx like HW4 did.")
img("p06_img1.png")

note(
    "a dedicated Create screenshot (the new course appearing in the list next to "
    "CreateRecord.jsx's dispatch(createCourse(...)) line) was captured earlier in this process "
    "but did not make it into this final screenshot set, a casualty of the time crunch finishing "
    "this homework. The Update and Delete screenshots below both show the create->update->delete "
    "lifecycle working end to end, just not a standalone Create frame."
)

h2("Redux: Update")
p("The course after being renamed to 'Agentic AI systems', shown next to UpdateRecord.jsx's "
  "dispatch(updateCourse(...)) line.")
img("p06_img2.png")

h2("Redux: Delete")
p("The delete confirmation screen for that same course, next to DeleteRecord.jsx's "
  "dispatch(deleteCourse(...)) line.")
img("p07_img1.png")

# ============================================================= PART 2
h1("Part 2: MCP Tool Servers")
p(
    "Two local MCP servers. meals_server.py wraps TheMealDB's public API with four tools, "
    "run over STDIO with FastMCP. domain_server.py is the required server over this project's "
    "own course catalogue data, with exactly three tools (search, detail lookup, aggregate), all "
    "returning one shared {ok, data, error} envelope that Part 4 reuses unchanged."
)

h2("meals server: search_meals_by_name")
p("Query 'Arrabiata', limit 3, run in the MCP Inspector.")
img("p08_img1.png")

note(
    "the Inspector's result panel was unreliable during this session (a known UI quirk also hit "
    "earlier in this homework), so labeled result screenshots for the other three meals tools "
    "(meals_by_ingredient, random_meal, meal_details) did not get captured before time ran out. "
    "All four tools were verified working directly in Python against the live TheMealDB API "
    "earlier in development (search for 'Arrabiata' returned one match, a chicken-ingredient "
    "filter returned three cards, a random meal returned a full recipe, and a lookup by a known "
    "id returned the correct meal name), and the message log on the right of this screenshot "
    "confirms the real STDIO tool call went through."
)

h2("domain server: search_courses (valid)")
p("Query 'cloud computing', the three matching courses returned in the Results panel.")
img("p09_img1.png")
img("p09_img2.png")

h2("domain server: search_courses (invalid) and course_detail (valid)")
note(
    "these two calls were set up in the Inspector (query left empty for search_courses; "
    "course_id 7 typed into course_detail) but the Execute Tool click and result capture did not "
    "happen before time ran out. Both were verified directly in Python earlier: an empty query "
    "returns {\"ok\": false, \"error\": \"query must be a non-empty string\"}, and course_detail(7) "
    "returns the full CS-00005 record with its Computer Science department attached (see "
    "METRICS.md's Part 2 table for the exact JSON)."
)
img("p10_img1.png")
img("p10_img2.png")

h2("domain server: course_detail (invalid)")
img("p11_img1.png")

h2("domain server: department_stats (valid)")
img("p11_img2.png")

h2("domain server: department_stats (invalid)")
img("p11_img3.png")

# ============================================================= PART 3
h1("Part 3: Tool Contracts Under Stress")
p(
    "Using the same three domain tools and the same rejected-call examples from Part 2B: the "
    "expected input schema, the rejected input, and the returned error for each is documented in "
    "METRICS.md section 3.18. call_with_retry wraps any callable with a per-attempt timeout and "
    "bounded exponential backoff. VERIFY_SEED makes the 150-call fault-injection experiment "
    "reproducible, confirmed by rerunning the same seed and rate and getting the identical "
    "outcome sequence both times."
)

h2("Retry demonstration and 150-call experiment")
p(
    "All three required scenarios (success on the first attempt, failure then success on retry, "
    "failure through every attempt raising a clean error instead of crashing), followed by the "
    "150-call summary at 0%, 20%, and 50% injected failure rates."
)
img("p12_img1.png")

t = doc.add_table(rows=0, cols=4)
t.style = "Table Grid"
for row_vals in [
    ("Injected failure rate", "Success rate", "Mean latency (ms)", "p99 latency (ms)"),
    ("0%", "1.00", "41.34", "54.55"),
    ("20%", "0.98", "53.43", "183.93"),
    ("50%", "0.82", "99.34", "186.00"),
]:
    row = t.add_row()
    for i, v in enumerate(row_vals):
        row.cells[i].text = v

p(
    "With max_retries=2 (3 attempts total) and each attempt independently drawing at the "
    "injected rate, the chance of a call exhausting every attempt is rate^3: 0% at a 0% injected "
    "rate, 0.8% at 20%, 12.5% at 50%. Over 50 calls that predicts roughly 0, 0.4, and 6.25 failed "
    "calls, close to what was observed. Latency rises with the injected rate because more calls "
    "need a second or third attempt, and each extra attempt adds its own work time plus a backoff "
    "sleep before it runs."
)

h2("Is this retry policy right for an interactive assistant?")
p(
    "Mostly yes, with one caveat. max_retries=2 and fast backoff (20ms, then 40ms) keep even the "
    "worst case here (p99 at 50%) under 200ms, well inside what feels instant in a chat UI. The "
    "real risk is the 1 second per-attempt timeout: a genuinely slow dependency could make three "
    "attempts take a few seconds combined, long enough for a user to assume the app froze. For an "
    "interactive assistant the timeout should probably drop to 300 to 500ms, fail fast and let "
    "the retry handle recovery instead of one long wait. For batch processing, where nobody is "
    "watching a spinner, I would raise max_retries to around 5, lengthen the backoff so retries "
    "are spaced further apart, and raise the timeout substantially, since a slow-but-eventually-"
    "successful response beats a wasted retry when nothing is time-sensitive."
)

# ============================================================= PART 4
h1("Part 4: One Safe Tool Entry Point and Offline Tests")
p(
    "execute_tool(name, inputs) is the single function the Part 5 agent calls to run any of the "
    "three domain tools. It reuses the same {ok, data, error} envelope, returns it as a JSON "
    "string, and never raises, even for an unknown tool name or malformed inputs. The offline "
    "test runner builds a temporary in-memory SQLite database from the same models the real "
    "MySQL database uses, so all ten tests run without touching the network, the real database, "
    "or an LLM."
)

h2("Offline test runner (10/10, includes the two Part 5 tests)")
img("p12_img2.png")

# ============================================================= PART 5
h1("Part 5: Agent Loop, Safety Rule, and Reflection")
p(
    "One business-logic safety rule lives inside execute_tool: agent calls to search_courses are "
    "capped at limit<=20, stricter than the 100 a trusted direct caller like Postman gets, so an "
    "autonomous loop cannot pull a large chunk of the catalogue into its own context in a single "
    "call. run_agent(user_input) drives a JSON-protocol tool-use loop against the real local "
    "Ollama model, with a turn counter and a max_steps ceiling, logging every step to "
    "agent_runs.jsonl."
)

h2("Safety rule: allowed call and blocked call")
p("A limit of 10 is allowed and returns real data; a limit of 50 is blocked by the safety rule "
  "before the search even runs.")
img("p12_img3.png")

h2("Four scenarios against the real model")
t = doc.add_table(rows=0, cols=4)
t.style = "Table Grid"
for row_vals in [
    ("#", "Steps", "Stop reason", "Tool calls"),
    ("1: What are the details for course id 7?", "2", "normal_completion", "1"),
    ("2: Give me the stats for department 1.", "2", "normal_completion", "1"),
    ("3: Search for courses about Cloud Computing.", "2", "normal_completion", "1"),
    ("4: Search for courses about data, with a limit of 50.", "4", "max_steps", "4"),
]:
    row = t.add_row()
    for i, v in enumerate(row_vals):
        row.cells[i].text = v
img("p13_img1.png")

h2("Reflection")
p(
    "Run picked: scenario 4, \"Search for courses about data, with a limit of 50.\" "
    "(stop_reason: max_steps, step_count: 4, tool_call_count: 4)."
)
p(
    "Step 1: the model called search_courses with query 'data' and limit 50, copying the limit "
    "straight from the user's wording. execute_tool's safety rule rejected it before the real "
    "search logic ever ran: \"safety rule: agent calls to search_courses are capped at limit<=20 "
    "(got 50)\". No exception, no crash, just a normal {ok: false, ...} envelope fed back into the "
    "conversation."
)
p(
    "Step 2 is where it gets interesting. Instead of retrying with a smaller limit, the correct "
    "fix and the number needed for it, 20, were right there in the error string, the model "
    "switched tools entirely and called department_stats with department_id set to the string "
    "'data', reusing the word from the user's question as if it were an id. That is a type error, "
    "not a safety violation, and domain_server.py's own input validation rejected it: "
    "\"department_id must be a positive integer\". Step 3 repeated the same mistake with "
    "course_detail(course_id='data'), rejected the same way. Step 4 went back to the original "
    "blocked call, search_courses with limit 50 again, unchanged. At that point max_steps=4 was "
    "reached and the harness stopped the loop with stop_reason max_steps and no final answer, "
    "rather than letting it spin indefinitely."
)
p(
    "What this shows: the safety rule worked exactly as designed, blocking the oversized call "
    "every time it was tried, including the repeat on step 4. The max_steps ceiling also worked "
    "as designed, turning what would otherwise be unbounded flailing into a bounded, clean stop. "
    "What did not work is the small model's own recovery behavior. A 1.5B-parameter model given a "
    "plain error string did not reliably parse \"capped at limit<=20\" into \"retry with limit=20\", "
    "and instead wandered into unrelated tools with nonsensical arguments. The safety and "
    "step-limit mechanisms are what kept a confused model from being a liability, they are not "
    "what makes the model smart about fixing its own mistakes, and this run is a clean, honest "
    "example of that gap."
)

# ============================================================= AI_USE
h1("AI Use Statement")

h2("1. What I used an AI assistant for, and what I did myself")
p(
    "I used Claude Code to design and write the Department entity and its migration against the "
    "already-populated HW4 database, extend the Course API with validation and proper error "
    "codes, migrate the React client's data layer to Redux Toolkit, build both MCP servers, write "
    "the retry/backoff policy and the fault-injection experiment, and build the Part 4 offline "
    "test runner and the Part 5 agent loop. It ran all of this against real infrastructure "
    "throughout: the actual MySQL container with the real HW4 dataset, the real TheMealDB API, "
    "the real MCP STDIO protocol, and the real local Ollama model."
)
p(
    "What I did myself: decided Department was the right related entity and that it should be a "
    "genuinely new table rather than repurposing HW4's free-text department column, which stays "
    "untouched since Part 3's N+1 seed data from HW4 depends on it. Chose the domain server's "
    "three tools and what the aggregate operation should mean for this domain. Designed the "
    "search_courses safety rule before any agent code existed to enforce it against. Picked which "
    "real agent run to walk through in the reflection, and read the raw agent_runs.jsonl trace "
    "myself to understand why it actually failed rather than just reporting the stop reason."
)

h2("2. One AI-produced output that was wrong or unsuitable")
p(
    "The first version of migrate_hw05.py deadlocked. It opened one SQLAlchemy session, used it "
    "to read back the newly-seeded Department rows, and then, without closing that read "
    "transaction, opened a second connection to run an ALTER TABLE statement adding a foreign key "
    "on courses referencing departments. Adding a foreign key needs a metadata lock on both "
    "tables involved, including the referenced one, and the first session's still-open read "
    "transaction was already holding a lock on departments that the DDL statement then had to "
    "wait on forever."
)

h2("3. How I detected the problem")
p(
    "The script sat with zero output for about three minutes. Rather than assume it was simply "
    "slow, I checked Get-Process CPU time for the two Python processes first: both were at a "
    "fraction of a second of CPU time after three minutes of wall time, which rules out slow but "
    "computing and points at blocked waiting on something. SHOW PROCESSLIST on the MySQL "
    "container confirmed it directly: one connection sitting in Sleep, a second stuck in Waiting "
    "for table metadata lock on the exact ALTER TABLE statement."
)

h2("4. What I changed and why it works now")
p(
    "I added an explicit commit right after each read that happens before a DDL step, closing out "
    "that transaction and its locks before the ALTER TABLE runs. After that fix the script "
    "completed in a couple of seconds, and rerunning it is idempotent, every step checks whether "
    "it already happened first, so I could verify it safely a second time. The broader lesson: "
    "with SQLAlchemy's default autocommit=False, even a plain SELECT opens a transaction that "
    "stays open, and holding one across a later schema change is enough to deadlock against "
    "yourself in a single script, not just against a separate concurrent process."
)

out_path = "Achutuni_HW5.docx"
doc.save(out_path)
print("wrote", out_path)
